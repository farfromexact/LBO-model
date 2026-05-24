from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from lbo.adapters.template_input_adapter import TemplateInputAdapter
from lbo.engine.template_memo_engine import calculate_memo_metrics
from lbo.io.template_generator import create_standard_template_bytes


def _model(years: int = 5):
    wb = load_workbook(BytesIO(create_standard_template_bytes(years)), data_only=True)
    return TemplateInputAdapter().extract(wb, f"{years}_year_template.xlsx")


@pytest.mark.parametrize("years", [3, 5, 7])
def test_standard_template_reads_valid_forecast_periods(years: int) -> None:
    model = _model(years)

    assert model.metadata.detected_template_id == "standard_input_template"
    assert len(model.financials.revenue) == years
    assert model.returns.moic.value is not None
    assert model.returns.irr.value is not None
    assert model.metrics["revenue"].source_sheet == "Operating Forecast"


def test_standard_template_rejects_missing_required_sheet() -> None:
    wb = load_workbook(BytesIO(create_standard_template_bytes()), data_only=True)
    del wb["Debt & Cash"]

    match = TemplateInputAdapter().can_handle(wb)

    assert "Debt & Cash" in match.missing_features
    with pytest.raises(ValueError):
        TemplateInputAdapter().extract(wb)


def test_standard_template_marks_missing_fields_without_zero_placeholder() -> None:
    wb = load_workbook(BytesIO(create_standard_template_bytes()), data_only=True)
    ws = wb["Deal Setup"]
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value == "sponsor_equity_invested":
            ws.cell(row, 3).value = None

    model = TemplateInputAdapter().extract(wb)

    assert "sponsor_equity_invested" in model.missing_fields
    assert model.returns.moic.value is None
    assert model.returns.irr.value is None
    assert model.returns.sponsor_equity_invested.value is None


def test_standard_template_core_return_math_and_memo_metrics() -> None:
    model = _model(5)
    memo = calculate_memo_metrics(model)

    assert model.returns.moic.value == pytest.approx(1040 / 400)
    assert model.returns.irr.value is not None
    assert memo.revenue_cagr is not None
    assert memo.ebitda_cagr is not None
    assert memo.value_creation_bridge.exit_equity_value == pytest.approx(model.valuation.exit_equity_value.value)
    assert model.debt.net_debt
    assert model.sensitivities


def test_standard_template_preserves_metric_lineage() -> None:
    model = _model(5)

    for metric in model.metrics.values():
        assert metric.source_sheet is not None
        assert metric.source_cell is not None or metric.source_range is not None

