from __future__ import annotations

from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import Workbook

from lbo.adapters import GenericExcelAdapter, detect_best_adapter
from lbo.engine import generate_sop_analysis, run_standardized_sensitivity


def generic_workbook_bytes() -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    financial = workbook.create_sheet("Financial Output")
    balance = workbook.create_sheet("Balance Sheet")
    cash_flow = workbook.create_sheet("Cash Flow")
    capex = workbook.create_sheet("Capex")
    returns = workbook.create_sheet("Return Model")

    for col, year in zip(["K", "L", "M", "N", "O", "P"], range(2025, 2031)):
        financial[f"{col}2"] = year
        capex[f"{col}2"] = year
    for col, year in zip(["P", "Q", "R", "S", "T"], range(2026, 2031)):
        balance[f"{col}2"] = year
        cash_flow[f"{col}2"] = year

    financial["B4"] = "Revenue"
    financial["K4"] = 1151.488690754277
    financial["L4"] = 1290.0069
    financial["M4"] = 1429.2017
    financial["N4"] = 1573.4187
    financial["O4"] = 1740.8288
    financial["P4"] = 1941.2009819748844
    financial["B96"] = "EBITDA"
    financial["K96"] = 131.26852833687414
    financial["L96"] = 154.3358
    financial["M96"] = 190.9088
    financial["N96"] = 225.5222
    financial["O96"] = 261.9912
    financial["P96"] = 302.50123053446566

    balance["B87"] = "Net Debt"
    balance["P87"] = -254.40379986948815
    balance["Q87"] = -212.49920600099574
    balance["R87"] = -151.68032766494616
    balance["S87"] = -40.61004012678342
    balance["T87"] = 106.96170401207013

    capex["B6"] = "Total Capex"
    capex["K6"] = 17089.000491903218
    capex["L6"] = 17759.10054109354
    capex["M6"] = 18496.210595202894
    capex["N6"] = 19307.031654723185
    capex["O6"] = 20198.934820195504
    capex["P6"] = 21180.028302215054

    cash_flow["B26"] = "FCF"
    cash_flow["P26"] = 10.026264529263258
    cash_flow["Q26"] = 28.818878336049565
    cash_flow["R26"] = 79.07028753816275
    cash_flow["S26"] = 115.57174413885355
    cash_flow["T26"] = 146.1074630147732

    returns["G19"] = 1130
    returns["G24"] = 695.7298970542984
    returns["R89"] = 3025.012305344657
    returns["R93"] = 3310.0814723715002
    returns["L101"] = -552.4331520544102
    for cell in ["M101", "N101", "O101", "P101", "Q101"]:
        returns[cell] = 0
    returns["R101"] = 1916.6349582337702
    for cell, year in zip(["L86", "M86", "N86", "O86", "P86", "Q86", "R86"], [2025, 2025, 2026, 2027, 2028, 2029, 2030]):
        returns[cell] = datetime(year, 12, 31)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_generic_adapter_detects_and_extracts_standard_model() -> None:
    workbook = generic_workbook_bytes()
    adapter = GenericExcelAdapter()
    match = detect_best_adapter(workbook, [adapter])

    assert match.can_handle
    assert match.confidence > 0.5

    model = adapter.extract(workbook)

    assert model.metadata.adapter_name == "generic_excel"
    assert model.metadata.extraction_coverage > 0.8
    assert model.metrics["revenue"].is_available
    assert model.valuation.entry_ev.value == pytest.approx(1130)
    assert model.returns.moic.value == pytest.approx(3.469442322760885)
    assert model.returns.irr.value == pytest.approx(0.2823092043399811)
    assert model.returns.sponsor_cash_flows


def test_sop_and_sensitivity_consume_standard_model() -> None:
    model = GenericExcelAdapter().extract(generic_workbook_bytes())
    sop = generate_sop_analysis(model)
    sensitivity = run_standardized_sensitivity(model)

    assert sop.deal_summary
    assert sop.data_quality
    assert sop.model_audit
    assert sensitivity.cases

