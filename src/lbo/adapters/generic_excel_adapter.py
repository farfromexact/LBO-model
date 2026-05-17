from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from lbo.adapters.base_adapter import AdapterMatch, BaseWorkbookAdapter
from lbo.schemas.standard_model import (
    FinancialStatement,
    ModelMetadata,
    ReturnSummary,
    StandardMetric,
    StandardModel,
    ValuationSummary,
)


class GenericExcelAdapter(BaseWorkbookAdapter):
    adapter_name = "generic_excel"

    def can_handle(self, workbook) -> AdapterMatch:
        try:
            wb = _load_workbook(workbook)
        except Exception as exc:
            return AdapterMatch(
                adapter_name=self.adapter_name,
                confidence=0,
                can_handle=False,
                detected_template="Unreadable workbook",
                warnings=[str(exc)],
            )

        sheet_names = wb.sheetnames
        score = 0.2
        reasons: list[str] = [f"{len(sheet_names)} sheets found"]
        if _find_sheet(wb, ["output", "财务", "income", "p&l"]):
            score += 0.2
            reasons.append("financial output sheet detected")
        if _find_sheet(wb, ["balance", "资产负债"]):
            score += 0.15
            reasons.append("balance sheet detected")
        if _find_sheet(wb, ["cash", "现金流"]):
            score += 0.15
            reasons.append("cash flow sheet detected")
        if _find_sheet(wb, ["return", "回报", "ev/ebitda", "evebitda"]):
            score += 0.2
            reasons.append("return model sheet detected")
        if _workbook_contains_any_label(wb, ["EBITDA", "IRR", "MOIC", "MOC", "企业价值"]):
            score += 0.1
            reasons.append("key LBO labels detected")

        wb.close()
        return AdapterMatch(
            adapter_name=self.adapter_name,
            confidence=min(score, 1.0),
            can_handle=score >= 0.35,
            detected_template="Generic Excel LBO / investment model",
            reasons=reasons,
        )

    def extract(self, workbook) -> StandardModel:
        wb = _load_workbook(workbook)
        workbook_name = _workbook_name(workbook)
        match = self.can_handle(workbook)

        financial_sheet = _find_sheet(wb, ["output_财务", "财务", "income", "p&l", "financial"])
        balance_sheet = _find_sheet(wb, ["资产负债", "balance"])
        cash_flow_sheet = _find_sheet(wb, ["现金流", "cash flow", "cf"])
        capex_sheet = _find_sheet(wb, ["capex", "产能"])
        return_sheet = _find_sheet(wb, ["回报", "return", "ev/ebitda", "evebitda"])

        metrics = {
            "revenue": _extract_series_metric(financial_sheet, "revenue", "Revenue", ["集团总营收", "营业收入", "Revenue"]),
            "ebitda": _extract_series_metric(financial_sheet, "ebitda", "EBITDA", ["EBITDA"]),
            "net_debt": _extract_series_metric(balance_sheet, "net_debt", "Net Debt", ["净债务", "Net Debt"]),
            "capex": _extract_series_metric(capex_sheet, "capex", "Capex", ["合计资本支出", "Total Capex"]),
            "cash_flow": _extract_series_metric(cash_flow_sheet, "cash_flow", "Cash Flow", ["净现金流", "经营性现金流", "Free Cash Flow", "FCF"]),
        }

        valuation = _extract_valuation(return_sheet, metrics)
        returns = _extract_returns(return_sheet)
        metrics.update(
            {
                "entry_ev": valuation.entry_ev,
                "entry_equity_value": valuation.entry_equity_value,
                "exit_ev": valuation.exit_ev,
                "exit_equity_value": valuation.exit_equity_value,
                "moic": returns.moic,
                "irr": returns.irr,
                "sponsor_equity_invested": returns.sponsor_equity_invested,
                "sponsor_proceeds": returns.sponsor_proceeds,
            }
        )

        required_ids = [
            "revenue",
            "ebitda",
            "net_debt",
            "capex",
            "cash_flow",
            "entry_ev",
            "entry_equity_value",
            "exit_ev",
            "exit_equity_value",
            "moic",
            "irr",
        ]
        missing_fields = [metric_id for metric_id in required_ids if not metrics[metric_id].is_available]
        warnings = list(returns.warnings)
        coverage = 0 if not required_ids else (len(required_ids) - len(missing_fields)) / len(required_ids)

        financials = {
            "operating": FinancialStatement(
                statement_id="operating",
                display_name="Operating Metrics",
                periods=_merge_periods(metrics["revenue"], metrics["ebitda"], metrics["cash_flow"]),
                metrics={key: metrics[key] for key in ["revenue", "ebitda", "cash_flow"]},
            )
        }

        wb.close()
        return StandardModel(
            metadata=ModelMetadata(
                source_workbook=workbook_name,
                adapter_name=self.adapter_name,
                adapter_confidence=match.confidence,
                detected_template=match.detected_template,
                extraction_coverage=coverage,
            ),
            metrics=metrics,
            financial_statements=financials,
            valuation=valuation,
            returns=returns,
            missing_fields=missing_fields,
            warnings=warnings,
        )


def _load_workbook(workbook) -> Workbook:
    source = workbook
    if isinstance(workbook, bytes):
        source = BytesIO(workbook)
    elif hasattr(workbook, "getvalue"):
        source = BytesIO(workbook.getvalue())
    return load_workbook(source, read_only=True, data_only=True)


def _workbook_name(workbook) -> str:
    if isinstance(workbook, (str, Path)):
        return Path(workbook).name
    return getattr(workbook, "name", "uploaded_workbook.xlsx")


def _find_sheet(wb: Workbook, tokens: list[str]) -> Worksheet | None:
    lowered = [(sheet_name, sheet_name.casefold().replace(" ", "").replace("_", "")) for sheet_name in wb.sheetnames]
    for sheet_name, normalized in lowered:
        for token in tokens:
            if token.casefold().replace(" ", "").replace("_", "") in normalized:
                return wb[sheet_name]
    return None


def _workbook_contains_any_label(wb: Workbook, labels: list[str]) -> bool:
    for ws in wb.worksheets:
        if _find_label_cell(ws, labels) is not None:
            return True
    return False


def _extract_series_metric(
    ws: Worksheet | None,
    metric_id: str,
    display_name: str,
    labels: list[str],
    unit: str | None = None,
) -> StandardMetric:
    if ws is None:
        return _missing_metric(metric_id, display_name, "required sheet not detected")

    label_cell = _find_label_cell(ws, labels)
    if label_cell is None:
        return _missing_metric(metric_id, display_name, "label not found")

    series: dict[str, float] = {}
    for column_index, cell in enumerate(ws[label_cell.row], start=1):
        if column_index <= label_cell.column or not isinstance(cell.value, (int, float)):
            continue
        period = _period_for_column(ws, column_index)
        if period is not None:
            series[period] = float(cell.value)

    if not series:
        return _missing_metric(metric_id, display_name, "no year-labeled numeric values found", ws.title, label_cell.coordinate)

    latest_period = sorted(series)[-1]
    return StandardMetric(
        metric_id=metric_id,
        display_name=display_name,
        value=series[latest_period],
        period=latest_period,
        unit=unit,
        source_sheet=ws.title,
        source_cell=label_cell.coordinate,
        confidence="medium",
        series=series,
    )


def _extract_valuation(return_sheet: Worksheet | None, metrics: dict[str, StandardMetric]) -> ValuationSummary:
    if return_sheet is None:
        missing = _missing_metric("valuation", "Valuation", "return model sheet not detected")
        return ValuationSummary(
            entry_ev=missing.model_copy(update={"metric_id": "entry_ev", "display_name": "Entry EV"}),
            entry_equity_value=missing.model_copy(update={"metric_id": "entry_equity_value", "display_name": "Entry Equity Value"}),
            exit_ev=missing.model_copy(update={"metric_id": "exit_ev", "display_name": "Exit EV"}),
            exit_equity_value=missing.model_copy(update={"metric_id": "exit_equity_value", "display_name": "Exit Equity Value"}),
            net_debt=metrics.get("net_debt", _missing_metric("net_debt", "Net Debt", "not extracted")),
        )

    return ValuationSummary(
        entry_ev=_metric_from_cell(return_sheet, "entry_ev", "Entry EV", "G19"),
        entry_equity_value=_metric_from_cell(return_sheet, "entry_equity_value", "Entry Equity Value", "G24"),
        exit_ev=_metric_from_cell(return_sheet, "exit_ev", "Exit EV", "R89"),
        exit_equity_value=_metric_from_cell(return_sheet, "exit_equity_value", "Exit Equity Value", "R93"),
        net_debt=metrics.get("net_debt", _missing_metric("net_debt", "Net Debt", "not extracted")),
    )


def _extract_returns(return_sheet: Worksheet | None) -> ReturnSummary:
    if return_sheet is None:
        missing = _missing_metric("returns", "Returns", "return model sheet not detected")
        return ReturnSummary(
            moic=missing.model_copy(update={"metric_id": "moic", "display_name": "MOIC"}),
            irr=missing.model_copy(update={"metric_id": "irr", "display_name": "IRR"}),
            sponsor_equity_invested=missing.model_copy(update={"metric_id": "sponsor_equity_invested", "display_name": "Sponsor Equity Invested"}),
            sponsor_proceeds=missing.model_copy(update={"metric_id": "sponsor_proceeds", "display_name": "Sponsor Proceeds"}),
            warnings=["Return model sheet not detected; IRR/MOIC not calculated."],
        )

    cash_flows = _sponsor_cash_flows(return_sheet)
    sponsor_equity = _metric_from_cell(return_sheet, "sponsor_equity_invested", "Sponsor Equity Invested", "L101")
    sponsor_proceeds = _metric_from_cell(return_sheet, "sponsor_proceeds", "Sponsor Proceeds", "R101")
    warnings: list[str] = []

    if len(cash_flows) < 2 or not any(cf["amount"] < 0 for cf in cash_flows) or not any(cf["amount"] > 0 for cf in cash_flows):
        warnings.append("Sponsor cash flow array missing or incomplete; IRR/MOIC not calculated.")
        moic = _missing_metric("moic", "MOIC", "sponsor cash flows missing")
        irr = _missing_metric("irr", "IRR", "sponsor cash flows missing")
    else:
        invested = abs(sum(cf["amount"] for cf in cash_flows if cf["amount"] < 0))
        proceeds = sum(cf["amount"] for cf in cash_flows if cf["amount"] > 0)
        moic_value = proceeds / invested if invested else None
        irr_value = _xirr(cash_flows)
        moic = StandardMetric(metric_id="moic", display_name="MOIC", value=moic_value, unit="x", confidence="high", source_sheet=return_sheet.title, source_cell="L101:R101")
        irr = StandardMetric(metric_id="irr", display_name="IRR", value=irr_value, unit="%", confidence="high", source_sheet=return_sheet.title, source_cell="L101:R101")

    return ReturnSummary(
        moic=moic,
        irr=irr,
        sponsor_equity_invested=sponsor_equity,
        sponsor_proceeds=sponsor_proceeds,
        sponsor_cash_flows=cash_flows,
        warnings=warnings,
    )


def _sponsor_cash_flows(ws: Worksheet) -> list[dict[str, Any]]:
    try:
        amounts = [cell.value for row in ws["L101:R101"] for cell in row]
        dates = [cell.value for row in ws["L86:R86"] for cell in row]
    except Exception:
        return []
    cash_flows = []
    for date_value, amount in zip(dates, amounts):
        if isinstance(amount, (int, float)):
            cash_flows.append({"date": _date(date_value), "amount": float(amount)})
    return cash_flows


def _xirr(cash_flows: list[dict[str, Any]]) -> float | None:
    dated_flows = [(cf["date"], cf["amount"]) for cf in cash_flows if cf.get("date") is not None]
    if len(dated_flows) < 2 or not any(amount < 0 for _, amount in dated_flows) or not any(amount > 0 for _, amount in dated_flows):
        return None

    start_date = dated_flows[0][0]

    def npv(rate: float) -> float:
        total = 0.0
        for flow_date, amount in dated_flows:
            years = (flow_date - start_date).days / 365.0
            total += amount / ((1 + rate) ** years)
        return total

    low = -0.999
    high = 10.0
    low_value = npv(low)
    high_value = npv(high)
    if low_value * high_value > 0:
        return None

    for _ in range(100):
        mid = (low + high) / 2
        mid_value = npv(mid)
        if abs(mid_value) < 1e-8:
            return mid
        if low_value * mid_value <= 0:
            high = mid
            high_value = mid_value
        else:
            low = mid
            low_value = mid_value
    return (low + high) / 2


def _metric_from_cell(ws: Worksheet, metric_id: str, display_name: str, cell_ref: str) -> StandardMetric:
    value = ws[cell_ref].value
    if not isinstance(value, (int, float)):
        return _missing_metric(metric_id, display_name, f"{cell_ref} is not numeric", ws.title, cell_ref)
    return StandardMetric(
        metric_id=metric_id,
        display_name=display_name,
        value=float(value),
        source_sheet=ws.title,
        source_cell=cell_ref,
        confidence="high",
    )


def _find_label_cell(ws: Worksheet, labels: list[str]) -> Cell | None:
    normalized_labels = [label.casefold() for label in labels]
    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str):
                continue
            value = cell.value.casefold().strip()
            if any(label in value for label in normalized_labels):
                return cell
    return None


def _period_for_column(ws: Worksheet, column: int) -> str | None:
    for row in range(1, min(ws.max_row, 8) + 1):
        value = ws.cell(row=row, column=column).value
        if isinstance(value, datetime):
            return str(value.year)
        if isinstance(value, int) and 2000 <= value <= 2100:
            return str(value)
        if isinstance(value, str):
            for year in range(2000, 2101):
                if str(year) in value:
                    return str(year)
    return None


def _date(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date()
    return value


def _missing_metric(
    metric_id: str,
    display_name: str,
    reason: str,
    source_sheet: str | None = None,
    source_cell: str | None = None,
) -> StandardMetric:
    return StandardMetric(
        metric_id=metric_id,
        display_name=display_name,
        source_sheet=source_sheet,
        source_cell=source_cell,
        confidence="missing",
        missing_reason=reason,
    )


def _merge_periods(*metrics: StandardMetric) -> list[str]:
    periods = set()
    for metric in metrics:
        periods.update(metric.series)
    return sorted(periods)
