from __future__ import annotations

from openpyxl.utils import get_column_letter

from lbo.adapters.base_adapter import (
    BaseWorkbookAdapter,
    build_metric,
    collect_excel_errors,
    get_openpyxl_workbook,
    row_value,
    sheet_match,
)
from lbo.schemas.standard_model import (
    DebtSummary,
    FinancialSeries,
    ModelMetadata,
    ReturnSummary,
    SensitivityMatrix,
    StandardModel,
    ValuationSummary,
)


RETURN_SHEET = "投资回报分析"
CONTROL_SHEET = "控制页"
SUMMARY_SHEET = "财务摘要"


class ChineseCompactLBOAdapter(BaseWorkbookAdapter):
    template_id = "chinese_compact_lbo"
    adapter_name = "Chinese Compact LBO Adapter"
    fingerprint_sheets = [
        CONTROL_SHEET,
        RETURN_SHEET,
        SUMMARY_SHEET,
        "利润表",
        "现金流量表",
        "资产负债表",
        "利润表构建",
        "负债, 资本支出及营运资本",
    ]

    def can_handle(self, workbook):
        return sheet_match(workbook, self.fingerprint_sheets, self.adapter_name, self.template_id)

    def extract(self, workbook, source_file: str = "uploaded_workbook.xlsx") -> StandardModel:
        wb = get_openpyxl_workbook(workbook)
        ret = wb[RETURN_SHEET]
        summary = wb[SUMMARY_SHEET]
        control = wb[CONTROL_SHEET] if CONTROL_SHEET in wb.sheetnames else None
        excel_errors = collect_excel_errors(wb)
        warnings = ["No single base case found; extracted scenario matrix only."]
        if excel_errors:
            warnings.append("Workbook contains Excel errors; valid values were still extracted.")

        financials = FinancialSeries(
            revenue=_summary_series(summary, 15, "revenue", "Revenue"),
            ebitda=_summary_series(summary, 26, "ebitda", "Normalized EBITDA"),
            net_income=_summary_series(summary, 29, "net_income", "Net Income"),
            net_debt=_summary_series(summary, 41, "net_debt", "Net Debt / Cash"),
            capex=_summary_series(summary, 59, "capex", "Capex"),
            fcf=_summary_series(summary, 62, "fcf", "Free Cash Flow"),
        )
        debt = DebtSummary(
            net_debt=_summary_series(summary, 41, "net_debt", "Net Debt / Cash"),
            net_debt_to_ebitda=_summary_series(summary, 44, "net_debt_to_ebitda", "Net Debt / EBITDA", unit="x"),
        )
        valuation = ValuationSummary(
            entry_date=row_value(ret, 5, 6, "entry_date", "Deal Close Date"),
            entry_ebitda=row_value(ret, 8, 6, "entry_ebitda", "2024E EBITDA", unit="EUR mm"),
            entry_multiple=row_value(ret, 9, 6, "entry_multiple", "Entry EV/EBITDA", unit="x"),
            entry_ev=row_value(ret, 10, 6, "entry_ev", "Entry EV", unit="EUR mm"),
            entry_equity_value=row_value(ret, 13, 6, "entry_equity_value", "Entry Equity Value", unit="EUR mm"),
        )
        returns = ReturnSummary(
            irr=None,
            moic=None,
            sponsor_cash_flows=_cash_flows(ret),
            sponsor_equity_invested=row_value(ret, 17, 6, "sponsor_equity_invested", "Sponsor Equity Invested", unit="EUR mm"),
        )
        metadata = ModelMetadata(
            source_file=source_file,
            detected_template_id=self.template_id,
            adapter_name=self.adapter_name,
            confidence_score=self.can_handle(wb).confidence_score,
            project_name=_project_name(control, ret),
            currency="EUR",
            unit="mm",
            warnings=warnings,
            missing_fields=["irr", "moic"],
            excel_errors=excel_errors,
        )
        return StandardModel(
            metadata=metadata,
            financials=financials,
            valuation=valuation,
            returns=returns,
            debt=debt,
            sensitivities=_exit_matrices(ret),
        )


def _project_name(control, ret) -> str | None:
    for sheet, cell in [(control, "B2"), (ret, "B1")]:
        if sheet is not None and sheet[cell].value:
            return str(sheet[cell].value)
    return None


def _summary_series(sheet, row: int, metric_id: str, label: str, unit: str = "EUR mm") -> dict[int, object]:
    result = {}
    for col in range(1, sheet.max_column + 1):
        year = sheet.cell(6, col).value
        value = sheet.cell(row, col).value
        if isinstance(year, (int, float)) and isinstance(value, (int, float)):
            result[int(year)] = build_metric(
                metric_id=metric_id,
                label=label,
                value=float(value),
                period=int(year),
                unit=unit,
                source_sheet=sheet.title,
                source_cell=sheet.cell(row, col).coordinate,
                extraction_method="fixed_cell",
                confidence="high",
            )
    return result


def _cash_flows(sheet) -> list[object]:
    flows = []
    for row in range(64, 75):
        for col in range(1, sheet.max_column + 1):
            value = sheet.cell(row, col).value
            if isinstance(value, (int, float)):
                flows.append(
                    build_metric(
                        metric_id="sponsor_cash_flow",
                        label=str(sheet.cell(row, 2).value or "Sponsor Cash Flow"),
                        value=float(value),
                        source_sheet=sheet.title,
                        source_cell=sheet.cell(row, col).coordinate,
                        extraction_method="fixed_cell",
                        confidence="medium",
                    )
                )
    return flows


def _exit_matrices(sheet) -> list[SensitivityMatrix]:
    blocks = [("2028 Exit", 8, 11), ("2029 Exit", 13, 16)]
    matrices: list[SensitivityMatrix] = []
    metrics = [("IRR", 61, "IRR"), ("MOIC", 60, "MOIC"), ("Exit Equity Value", 52, "Exit Equity Value")]
    for metric_name, row, title_metric in metrics:
        row_values = [block[0] for block in blocks]
        col_values = []
        for col in range(blocks[0][1], blocks[0][2] + 1):
            col_values.append(sheet.cell(47, col).value)
        values = []
        for _, start_col, end_col in blocks:
            values.append([
                float(sheet.cell(row, col).value) if isinstance(sheet.cell(row, col).value, (int, float)) else None
                for col in range(start_col, end_col + 1)
            ])
        matrices.append(
            SensitivityMatrix(
                matrix_id=f"spark_{metric_name.lower().replace(' ', '_')}",
                title=f"{title_metric} by exit date and exit multiple",
                row_axis_name="Exit Date",
                col_axis_name="Exit EV/EBITDA Multiple",
                row_values=row_values,
                col_values=col_values,
                values=values,
                metric=metric_name,
                source_sheet=sheet.title,
                source_range=f"H{row}:P{row}",
                confidence="high",
            )
        )
    return matrices
