from __future__ import annotations

from openpyxl.utils import get_column_letter

from lbo.adapters.base_adapter import (
    AdapterMatch,
    BaseWorkbookAdapter,
    build_metric,
    collect_excel_errors,
    get_openpyxl_workbook,
    row_value,
    sheet_match,
)
from lbo.schemas.standard_model import (
    DebtSummary,
    FinancialSeries,
    ModelMetadata,
    ReturnSummary,
    SensitivityMatrix,
    StandardModel,
    ValuationSummary,
)


RETURN_SHEET = "回报测算-EVEBITDA（核心假设）"


class DragonDetailedLBOAdapter(BaseWorkbookAdapter):
    template_id = "dragon_detailed_lbo"
    adapter_name = "Dragon Detailed LBO Adapter"
    fingerprint_sheets = [
        RETURN_SHEET,
        "Output_财务",
        "Output_产能",
        "Output_区域剂型毛利",
        "利润表",
        "资产负债表",
        "现金流量表",
        "Ratio",
        "债务&权益",
        "贷款明细",
    ]

    def can_handle(self, workbook) -> AdapterMatch:
        return sheet_match(workbook, self.fingerprint_sheets, self.adapter_name, self.template_id)

    def extract(self, workbook, source_file: str = "uploaded_workbook.xlsx") -> StandardModel:
        wb = get_openpyxl_workbook(workbook)
        sheet = wb[RETURN_SHEET]
        warnings: list[str] = []
        excel_errors = collect_excel_errors(wb)
        if excel_errors:
            warnings.append("Workbook contains Excel errors; valid values were still extracted.")

        financials = FinancialSeries(
            revenue=_series(_sheet(wb, "Output_财务"), 4, "revenue", "Revenue"),
            ebitda=_series(_sheet(wb, "Output_财务"), 96, "ebitda", "EBITDA"),
            capex=_series(_sheet(wb, "Output_产能"), 6, "capex", "Capex"),
            fcf=_series(_sheet(wb, "现金流量表"), 26, "fcf", "Cash Flow"),
            net_debt=_series(_sheet(wb, "资产负债表"), 87, "net_debt", "Net Debt"),
        )
        valuation = ValuationSummary(
            entry_ev=row_value(sheet, 19, 7, "entry_ev", "Entry EV", unit="RMB mm"),
            entry_ebitda=row_value(sheet, 20, 7, "entry_ebitda", "Entry EBITDA", unit="RMB mm"),
            entry_multiple=row_value(sheet, 21, 7, "entry_multiple", "Entry EV/EBITDA", unit="x"),
            entry_equity_value=row_value(sheet, 24, 7, "entry_equity_value", "Entry Equity Value", unit="RMB mm"),
            exit_multiple=row_value(sheet, 87, 18, "exit_multiple", "Exit EV/EBITDA", unit="x"),
            exit_date=row_value(sheet, 86, 18, "exit_date", "Exit Date"),
            exit_ebitda=row_value(sheet, 88, 18, "exit_ebitda", "Exit EBITDA", unit="RMB mm"),
            exit_ev=row_value(sheet, 89, 18, "exit_ev", "Exit EV", unit="RMB mm"),
            exit_equity_value=row_value(sheet, 93, 18, "exit_equity_value", "Exit Equity Value", unit="RMB mm"),
        )
        sponsor_cash_flows = []
        for col in range(12, 19):
            cell = sheet.cell(101, col)
            date_cell = sheet.cell(86, col)
            if isinstance(cell.value, (int, float)):
                sponsor_cash_flows.append(
                    build_metric(
                        metric_id="sponsor_cash_flow",
                        label="Investor Net Cash Flow",
                        value=float(cell.value),
                        period=date_cell.value,
                        unit="RMB mm",
                        source_sheet=sheet.title,
                        source_cell=cell.coordinate,
                        extraction_method="fixed_cell",
                        confidence="high",
                    )
                )
        returns = ReturnSummary(
            sponsor_cash_flows=sponsor_cash_flows,
            sponsor_equity_invested=row_value(sheet, 55, 7, "sponsor_equity_invested", "Sponsor Equity Invested", unit="RMB mm"),
            sponsor_exit_proceeds=row_value(sheet, 101, 18, "sponsor_exit_proceeds", "Sponsor Exit Proceeds", unit="RMB mm"),
            moic=row_value(sheet, 103, 7, "moic", "MOIC", unit="x"),
            irr=row_value(sheet, 104, 7, "irr", "IRR", unit="%"),
        )
        sensitivities = _extract_sensitivities(sheet)
        metadata = ModelMetadata(
            source_file=source_file,
            detected_template_id=self.template_id,
            adapter_name=self.adapter_name,
            confidence_score=self.can_handle(wb).confidence_score,
            project_name=str(sheet["A1"].value) if sheet["A1"].value else None,
            currency="RMB",
            unit="mm",
            warnings=warnings,
            missing_fields=[],
            excel_errors=excel_errors,
        )
        return StandardModel(
            metadata=metadata,
            financials=financials,
            valuation=valuation,
            returns=returns,
            debt=DebtSummary(net_debt=financials.net_debt or {}),
            sensitivities=sensitivities,
        )


def _series(sheet, row: int, metric_id: str, label: str) -> dict[int, object]:
    if sheet is None:
        return {}
    result = {}
    header_row = _best_year_header_row(sheet, row)
    for col in range(1, sheet.max_column + 1):
        year = sheet.cell(header_row, col).value
        value = sheet.cell(row, col).value
        if hasattr(year, "year"):
            year = year.year
        if isinstance(year, (int, float)) and isinstance(value, (int, float)) and 1900 <= int(year) <= 2100:
            result[int(year)] = build_metric(
                metric_id=metric_id,
                label=label,
                value=float(value),
                period=int(year),
                unit="RMB mm",
                source_sheet=sheet.title,
                source_cell=sheet.cell(row, col).coordinate,
                extraction_method="fixed_cell",
                confidence="high",
            )
    return result


def _sheet(wb, name: str):
    return wb[name] if name in wb.sheetnames else None


def _best_year_header_row(sheet, data_row: int) -> int:
    candidate_rows = [2, 3, 5, max(1, data_row - 1)]
    best_row = candidate_rows[0]
    best_count = -1
    for row in candidate_rows:
        count = 0
        for col in range(1, sheet.max_column + 1):
            value = sheet.cell(row, col).value
            year = value.year if hasattr(value, "year") else value
            if isinstance(year, (int, float)) and 1900 <= int(year) <= 2100:
                count += 1
        if count > best_count:
            best_count = count
            best_row = row
    return best_row


def _extract_sensitivities(sheet) -> list[SensitivityMatrix]:
    matrices: list[SensitivityMatrix] = []
    for idx, start_row in enumerate([127, 149, 171], start=1):
        row_values = []
        col_values = []
        values = []
        for col in range(8, 14):
            value = sheet.cell(start_row + 2, col).value
            if isinstance(value, (int, float, str)) and value != "":
                col_values.append(value)
        for row in range(start_row + 3, min(start_row + 12, sheet.max_row) + 1):
            row_axis = sheet.cell(row, 7).value
            if row_axis is None:
                continue
            row_values.append(row_axis)
            values.append([
                float(sheet.cell(row, col).value) if isinstance(sheet.cell(row, col).value, (int, float)) else None
                for col in range(8, 8 + len(col_values))
            ])
        if row_values and col_values:
            matrices.append(
                SensitivityMatrix(
                    matrix_id=f"dragon_sensitivity_{idx}",
                    title=str(sheet.cell(start_row, 2).value or f"Sensitivity {idx}"),
                    row_axis_name="Entry Multiple",
                    col_axis_name="Exit Multiple",
                    row_values=row_values,
                    col_values=col_values,
                    values=values,
                    metric="IRR",
                    source_sheet=sheet.title,
                    source_range=f"G{start_row + 2}:{get_column_letter(7 + len(col_values))}{start_row + 2 + len(row_values)}",
                    confidence="medium",
                )
            )
    return matrices
