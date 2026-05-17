from __future__ import annotations

import re
from abc import ABC, abstractmethod
from datetime import date, datetime
from typing import Any, Iterable

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, Field

from lbo.schemas.standard_model import MetricValue, StandardModel


class AdapterMatch(BaseModel):
    template_id: str
    adapter_name: str
    confidence_score: float = Field(ge=0, le=1)
    matched_features: list[str] = Field(default_factory=list)
    missing_features: list[str] = Field(default_factory=list)
    warnings: list[str] = Field(default_factory=list)

    @property
    def confidence(self) -> float:
        return self.confidence_score

    @property
    def can_handle(self) -> bool:
        if self.template_id == "generic_lbo":
            return self.confidence_score >= 0.40
        return self.confidence_score >= 0.60

    @property
    def detected_template(self) -> str:
        return self.template_id

    @property
    def reasons(self) -> list[str]:
        return self.matched_features


class BaseWorkbookAdapter(ABC):
    template_id: str
    adapter_name: str

    @abstractmethod
    def can_handle(self, workbook) -> AdapterMatch:
        raise NotImplementedError

    @abstractmethod
    def extract(self, workbook, source_file: str) -> StandardModel:
        raise NotImplementedError


def get_openpyxl_workbook(workbook):
    return getattr(workbook, "workbook", workbook)


def normalize_label(text: Any) -> str:
    if text is None:
        return ""
    text = str(text).strip().lower()
    return re.sub(r"[\s\u3000:：()（）/\\_\-]+", "", text)


def find_sheet(workbook, possible_names: Iterable[str]) -> Worksheet | None:
    wb = get_openpyxl_workbook(workbook)
    normalized = {normalize_label(name): name for name in wb.sheetnames}
    for candidate in possible_names:
        key = normalize_label(candidate)
        if key in normalized:
            return wb[normalized[key]]
    for candidate in possible_names:
        key = normalize_label(candidate)
        for normalized_name, actual_name in normalized.items():
            if key and (key in normalized_name or normalized_name in key):
                return wb[actual_name]
    return None


def find_label_cell(sheet: Worksheet, label_patterns: Iterable[str], max_rows: int | None = None, max_cols: int | None = None):
    patterns = [normalize_label(pattern) for pattern in label_patterns]
    max_row = min(sheet.max_row, max_rows or sheet.max_row)
    max_col = min(sheet.max_column, max_cols or sheet.max_column)
    for row in sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            label = normalize_label(cell.value)
            if label and any(pattern and pattern in label for pattern in patterns):
                return cell
    return None


def detect_year_header_row(sheet: Worksheet, max_rows: int = 20) -> int | None:
    best_row: int | None = None
    best_count = 0
    for row_idx in range(1, min(sheet.max_row, max_rows) + 1):
        count = 0
        for cell in sheet[row_idx]:
            year = _coerce_year(cell.value)
            if year is not None:
                count += 1
        if count > best_count:
            best_count = count
            best_row = row_idx
    return best_row if best_count >= 2 else None


def extract_row_series(sheet: Worksheet, label_patterns: Iterable[str], header_row: int | None = None) -> dict[int, MetricValue]:
    label_cell = find_label_cell(sheet, label_patterns)
    if label_cell is None:
        return {}
    header = header_row or detect_year_header_row(sheet) or max(1, label_cell.row - 1)
    result: dict[int, MetricValue] = {}
    for col in range(label_cell.column + 1, sheet.max_column + 1):
        year = _coerce_year(sheet.cell(header, col).value)
        value = _numeric(sheet.cell(label_cell.row, col).value)
        if year is None or value is None:
            continue
        result[year] = build_metric(
            metric_id=normalize_label(str(label_patterns[0])) or str(label_cell.value),
            label=str(label_cell.value),
            value=value,
            period=year,
            source_sheet=sheet.title,
            source_cell=sheet.cell(label_cell.row, col).coordinate,
            extraction_method="row_series",
            confidence="medium",
        )
    return result


def get_value_right_of_label(sheet: Worksheet, label_patterns: Iterable[str], offset_cols: int = 1):
    label_cell = find_label_cell(sheet, label_patterns)
    if label_cell is None:
        return None, None
    value_cell = sheet.cell(label_cell.row, label_cell.column + offset_cols)
    return value_cell.value, value_cell.coordinate


def build_metric(
    metric_id: str,
    label: str,
    value: Any = None,
    period: int | str | None = None,
    unit: str | None = None,
    scale: float | None = 1.0,
    sign_convention: str | None = None,
    source_sheet: str | None = None,
    source_cell: str | None = None,
    source_range: str | None = None,
    extraction_method: str = "fixed_cell",
    confidence: str = "high",
    warning: str | None = None,
) -> MetricValue:
    if value is not None and isinstance(value, datetime):
        value = value.date()
    if isinstance(period, datetime):
        period = period.date().isoformat()
    elif isinstance(period, date):
        period = period.isoformat()
    if value is not None and not isinstance(value, (float, int, str, date)):
        value = str(value)
    if isinstance(value, int):
        value = float(value)
    return MetricValue(
        metric_id=metric_id,
        label=label,
        value=value,
        period=period,
        unit=unit,
        scale=scale,
        sign_convention=sign_convention,
        source_sheet=source_sheet,
        source_cell=source_cell,
        source_range=source_range,
        extraction_method=extraction_method,
        confidence=confidence,
        warning=warning,
    )


def missing_metric(metric_id: str, label: str, warning: str) -> MetricValue:
    return build_metric(metric_id, label, None, extraction_method="missing", confidence="missing", warning=warning)


def collect_excel_errors(workbook) -> list[str]:
    wb = get_openpyxl_workbook(workbook)
    errors: list[str] = []
    error_values = {"#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NUM!", "#NULL!"}
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip() in error_values:
                    errors.append(f"{sheet.title}!{cell.coordinate}: {cell.value.strip()}")
    return errors


def value_at(sheet: Worksheet, cell_ref: str, metric_id: str, label: str, **kwargs) -> MetricValue:
    return build_metric(
        metric_id=metric_id,
        label=label,
        value=sheet[cell_ref].value,
        source_sheet=sheet.title,
        source_cell=cell_ref,
        extraction_method="fixed_cell",
        **kwargs,
    )


def row_value(sheet: Worksheet, row: int, col: int, metric_id: str, label: str, **kwargs) -> MetricValue:
    cell = sheet.cell(row, col)
    return build_metric(
        metric_id=metric_id,
        label=label,
        value=cell.value,
        source_sheet=sheet.title,
        source_cell=cell.coordinate,
        extraction_method="fixed_cell",
        **kwargs,
    )


def sheet_match(workbook, sheet_names: list[str], adapter_name: str, template_id: str) -> AdapterMatch:
    wb = get_openpyxl_workbook(workbook)
    normalized_existing = {normalize_label(name) for name in wb.sheetnames}
    matched = [name for name in sheet_names if normalize_label(name) in normalized_existing]
    missing = [name for name in sheet_names if normalize_label(name) not in normalized_existing]
    score = len(matched) / len(sheet_names) if sheet_names else 0.0
    return AdapterMatch(
        template_id=template_id,
        adapter_name=adapter_name,
        confidence_score=score,
        matched_features=matched,
        missing_features=missing,
    )


def _coerce_year(value: Any) -> int | None:
    if isinstance(value, datetime):
        return value.year
    if isinstance(value, date):
        return value.year
    if isinstance(value, (int, float)) and 1900 <= int(value) <= 2100:
        return int(value)
    if isinstance(value, str):
        match = re.search(r"(19|20)\d{2}", value)
        if match:
            return int(match.group(0))
    return None


def _numeric(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    return None
