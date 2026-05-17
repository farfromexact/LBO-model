from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from lbo.adapters.base_adapter import (
    AdapterMatch,
    BaseWorkbookAdapter,
    build_metric,
    collect_excel_errors,
    extract_row_series,
    find_label_cell,
    get_openpyxl_workbook,
    missing_metric,
)
from lbo.schemas.standard_model import FinancialSeries, ModelMetadata, ReturnSummary, StandardModel, ValuationSummary


class GenericLBOAdapter(BaseWorkbookAdapter):
    template_id = "generic_lbo"
    adapter_name = "generic_lbo"

    def can_handle(self, workbook) -> AdapterMatch:
        wb = _wb(workbook)
        labels = _scan_labels(wb)
        matched: list[str] = []
        for feature, patterns in {
            "return sheet labels": ["irr", "moic", "moc", "投资回报", "回报测算", "return"],
            "financial labels": ["revenue", "收入", "ebitda", "fcf", "free cash flow", "capex", "资本支出"],
            "debt labels": ["net debt", "净债务", "债务"],
        }.items():
            if any(pattern.lower() in labels for pattern in patterns):
                matched.append(feature)
        confidence = min(0.59, 0.25 + 0.15 * len(matched))
        return AdapterMatch(
            template_id=self.template_id,
            adapter_name=self.adapter_name,
            confidence_score=confidence,
            matched_features=matched,
            missing_features=[],
            warnings=["Generic fallback uses low/medium confidence label matching only."],
        )

    def extract(self, workbook, source_file: str = "uploaded_workbook.xlsx") -> StandardModel:
        wb = _wb(workbook)
        return_sheet = _find_best_sheet(wb, ["irr", "moic", "moc", "投资回报", "回报测算", "return"])
        financial_sheet = _find_best_sheet(wb, ["revenue", "收入", "ebitda", "fcf", "free cash flow", "capex", "资本支出"])
        capex_sheet = _find_best_sheet(wb, ["capex", "资本支出"])
        cash_flow_sheet = _find_best_sheet(wb, ["fcf", "free cash flow", "cash flow", "现金流"])
        debt_sheet = _find_best_sheet(wb, ["net debt", "净债务", "债务"])
        warnings = ["Generic adapter selected; values are not high confidence."]

        financials = FinancialSeries()
        if financial_sheet is not None:
            financials.revenue = _series(financial_sheet, ["Revenue", "收入"], "revenue", "Revenue")
            financials.ebitda = _series(financial_sheet, ["EBITDA", "正常化EBITDA"], "ebitda", "EBITDA")
            financials.capex = _series(capex_sheet or financial_sheet, ["Total Capex", "Capex", "资本支出"], "capex", "Capex")
            financials.fcf = _series(cash_flow_sheet or financial_sheet, ["FCF", "Free Cash Flow", "自由现金流"], "fcf", "FCF")
        if debt_sheet is not None:
            financials.net_debt = _series(debt_sheet, ["Net Debt", "净债务", "净负债"], "net_debt", "Net Debt")

        valuation = ValuationSummary(
            entry_ev=_metric_near(return_sheet, ["Entry EV", "Enterprise Value", "企业价值"], "entry_ev", "Entry EV"),
            exit_ev=_metric_near(return_sheet, ["Exit EV", "Enterprise Value at Exit", "退出企业价值"], "exit_ev", "Exit EV"),
        )
        returns = ReturnSummary(
            irr=_metric_near(return_sheet, ["IRR", "净IRR"], "irr", "IRR", unit="%"),
            moic=_metric_near(return_sheet, ["MOIC", "MoC", "MOC"], "moic", "MOIC", unit="x"),
        )
        if return_sheet is not None and return_sheet.title == "Return Model":
            valuation, returns = _extract_simple_return_model(return_sheet, valuation, returns)
        metadata = ModelMetadata(
            source_file=source_file,
            detected_template_id=self.template_id,
            adapter_name=self.adapter_name,
            confidence_score=0.40,
            warnings=warnings,
            missing_fields=[],
            excel_errors=collect_excel_errors(wb),
        )
        return StandardModel(metadata=metadata, financials=financials, valuation=valuation, returns=returns)


class GenericExcelAdapter(GenericLBOAdapter):
    adapter_name = "generic_excel"


def _wb(workbook):
    if isinstance(workbook, bytes):
        return load_workbook(BytesIO(workbook), data_only=True)
    return get_openpyxl_workbook(workbook)


def _scan_labels(wb) -> str:
    values: list[str] = []
    for sheet in wb.worksheets:
        values.append(sheet.title.lower())
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 80), min_col=1, max_col=min(sheet.max_column, 20)):
            for cell in row:
                if cell.value is not None:
                    values.append(str(cell.value).lower())
    return " ".join(values)


def _find_best_sheet(wb, patterns: list[str]):
    best = None
    best_score = 0
    for sheet in wb.worksheets:
        text = sheet.title.lower()
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 120), min_col=1, max_col=min(sheet.max_column, 20)):
            for cell in row:
                if cell.value is not None:
                    text += " " + str(cell.value).lower()
        score = sum(1 for pattern in patterns if pattern.lower() in text)
        if score > best_score:
            best = sheet
            best_score = score
    return best


def _series(sheet, labels: list[str], metric_id: str, label: str):
    series = extract_row_series(sheet, labels)
    return {
        period: metric.model_copy(update={"metric_id": metric_id, "label": label, "confidence": "medium"})
        for period, metric in series.items()
    }


def _metric_near(sheet, labels: list[str], metric_id: str, label: str, unit: str | None = None):
    if sheet is None:
        return missing_metric(metric_id, label, "No likely return sheet found")
    label_cell = find_label_cell(sheet, labels, max_rows=120, max_cols=30)
    if label_cell is None:
        return missing_metric(metric_id, label, "Label not found by generic adapter")
    for col in range(label_cell.column + 1, min(sheet.max_column, label_cell.column + 12) + 1):
        value = sheet.cell(label_cell.row, col).value
        if isinstance(value, (int, float)):
            return build_metric(
                metric_id=metric_id,
                label=label,
                value=float(value),
                unit=unit,
                source_sheet=sheet.title,
                source_cell=sheet.cell(label_cell.row, col).coordinate,
                extraction_method="label_match",
                confidence="low",
            )
    return missing_metric(metric_id, label, "No numeric value found near label")


def _extract_simple_return_model(sheet, valuation: ValuationSummary, returns: ReturnSummary):
    import pyxirr

    valuation.entry_ev = build_metric("entry_ev", "Entry EV", sheet["G19"].value, unit="RMB mm", source_sheet=sheet.title, source_cell="G19", confidence="medium")
    valuation.entry_equity_value = build_metric(
        "entry_equity_value", "Entry Equity Value", sheet["G24"].value, unit="RMB mm", source_sheet=sheet.title, source_cell="G24", confidence="medium"
    )
    valuation.exit_ev = build_metric("exit_ev", "Exit EV", sheet["R89"].value, unit="RMB mm", source_sheet=sheet.title, source_cell="R89", confidence="medium")
    valuation.exit_equity_value = build_metric(
        "exit_equity_value", "Exit Equity Value", sheet["R93"].value, unit="RMB mm", source_sheet=sheet.title, source_cell="R93", confidence="medium"
    )
    cash_flows = []
    dates = []
    for col in range(12, 19):
        value = sheet.cell(101, col).value
        date_value = sheet.cell(86, col).value
        if isinstance(value, (int, float)):
            cash_flows.append(float(value))
            dates.append(date_value)
            returns.sponsor_cash_flows.append(
                build_metric(
                    "sponsor_cash_flow",
                    "Sponsor Cash Flow",
                    float(value),
                    period=date_value,
                    unit="RMB mm",
                    source_sheet=sheet.title,
                    source_cell=sheet.cell(101, col).coordinate,
                    confidence="medium",
                )
            )
    if cash_flows:
        invested = abs(cash_flows[0])
        proceeds = sum(value for value in cash_flows[1:] if value > 0)
        returns.sponsor_equity_invested = build_metric(
            "sponsor_equity_invested", "Sponsor Equity Invested", invested, unit="RMB mm", source_sheet=sheet.title, source_cell="L101", confidence="medium"
        )
        returns.sponsor_exit_proceeds = build_metric(
            "sponsor_exit_proceeds", "Sponsor Exit Proceeds", proceeds, unit="RMB mm", source_sheet=sheet.title, source_cell="R101", confidence="medium"
        )
        returns.moic = build_metric("moic", "MOIC", proceeds / invested if invested else None, unit="x", source_sheet=sheet.title, source_cell="L101:R101", confidence="medium")
        try:
            irr = pyxirr.xirr(dates, cash_flows)
        except Exception:
            irr = None
        returns.irr = build_metric("irr", "IRR", irr, unit="%", source_sheet=sheet.title, source_cell="L101:R101", confidence="medium")
    return valuation, returns
