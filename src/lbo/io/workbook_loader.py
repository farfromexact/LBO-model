from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook as openpyxl_load_workbook
from openpyxl.utils.exceptions import InvalidFileException


class WorkbookLoadError(Exception):
    def __init__(self, code: str, message: str) -> None:
        self.code = code
        self.message = message
        super().__init__(message)


@dataclass
class LoadedWorkbook:
    workbook: Any
    source_file: str
    warnings: list[str] = field(default_factory=list)
    encrypted: bool = False


def load_workbook(path, password_map: dict[str, str] | None = None, source_file: str | None = None) -> LoadedWorkbook:
    password_map = password_map or {}
    raw = _read_bytes(path)
    source = source_file or _source_name(path)

    try:
        wb = openpyxl_load_workbook(BytesIO(raw), data_only=True)
        return LoadedWorkbook(workbook=wb, source_file=source)
    except Exception as exc:
        if not _looks_encrypted(raw, exc):
            raise WorkbookLoadError("unreadable_workbook", f"Workbook could not be opened: {exc}") from exc

    password = _find_password(source, password_map)
    if not password:
        raise WorkbookLoadError(
            "encrypted_password_required",
            "Workbook appears to be encrypted and requires a password. Provide it through password_map.",
        )

    try:
        import msoffcrypto
    except ImportError as exc:
        raise WorkbookLoadError(
            "missing_dependency",
            "Encrypted workbook support requires msoffcrypto-tool.",
        ) from exc

    decrypted = BytesIO()
    try:
        office_file = msoffcrypto.OfficeFile(BytesIO(raw))
        office_file.load_key(password=password)
        office_file.decrypt(decrypted)
        decrypted.seek(0)
        wb = openpyxl_load_workbook(decrypted, data_only=True)
        return LoadedWorkbook(workbook=wb, source_file=source, encrypted=True)
    except Exception as exc:
        raise WorkbookLoadError("encrypted_open_failed", f"Encrypted workbook could not be decrypted: {exc}") from exc


def _read_bytes(path) -> bytes:
    if isinstance(path, bytes):
        return path
    if hasattr(path, "read"):
        position = path.tell() if hasattr(path, "tell") else None
        data = path.read()
        if position is not None and hasattr(path, "seek"):
            path.seek(position)
        return data
    return Path(path).read_bytes()


def _source_name(path) -> str:
    if isinstance(path, (str, Path)):
        return Path(path).name
    return getattr(path, "name", "uploaded_workbook.xlsx")


def _find_password(source_file: str, password_map: dict[str, str]) -> str | None:
    lowered = source_file.lower()
    for key, password in password_map.items():
        if key.lower() in lowered:
            return password
    return None


def _looks_encrypted(raw: bytes, exc: Exception) -> bool:
    if raw[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return True
    if isinstance(exc, (InvalidFileException,)):
        return True
    message = str(exc).lower()
    return "not a zip file" in message or "encrypted" in message or "ole" in message
