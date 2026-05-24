from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


TEMPLATE_SHEETS = ["Deal Setup", "Operating Forecast", "Debt & Cash", "Return Assumptions", "Qualitative Inputs"]


DEAL_FIELDS = [
    ("project_name", "Project Name", "Example Deal"),
    ("company", "Company", "ExampleCo"),
    ("sector", "Sector", "Business Services"),
    ("currency", "Currency", "USD"),
    ("unit", "Unit", "mm"),
    ("entry_date", "Entry Date", "2026-12-31"),
    ("exit_date", "Exit Date", "2031-12-31"),
    ("entry_ev", "Entry EV", 1000.0),
    ("entry_equity_value", "Entry Equity Value", 600.0),
    ("sponsor_equity_invested", "Sponsor Equity Invested", 400.0),
    ("sponsor_ownership", "Sponsor Ownership", 0.67),
    ("entry_multiple", "Entry EV / EBITDA", 10.0),
    ("exit_multiple", "Exit EV / EBITDA", 11.0),
    ("transaction_fees", "Transaction Fees", 20.0),
    ("rollover_management_equity", "Rollover / Management Equity", 0.0),
]

RETURN_FIELDS = [
    ("exit_ebitda", "Exit EBITDA", 160.0),
    ("exit_ev", "Exit EV", 1760.0),
    ("exit_net_debt", "Exit Net Debt", 200.0),
    ("exit_equity_value", "Exit Equity Value", 1560.0),
    ("dividends_cash_distributions", "Dividends / Cash Distributions", 0.0),
    ("sponsor_proceeds", "Sponsor Proceeds", 1040.0),
    ("moic", "MOIC", None),
    ("irr", "IRR", None),
]

QUAL_FIELDS = [
    ("investment_thesis", "Investment Thesis", "Why this deal is attractive."),
    ("key_risks", "Key Risks", "Main risks to underwrite."),
    ("diligence_questions", "Diligence Questions", "Open diligence questions."),
    ("management_market_notes", "Management / Market Notes", "Notes from management and market work."),
    ("exit_rationale", "Exit Rationale", "Likely exit path and buyer universe."),
]


def create_standard_template_bytes(years: int = 5) -> bytes:
    if not 3 <= years <= 7:
        raise ValueError("Standard template forecast period must be 3-7 years")

    wb = Workbook()
    wb.remove(wb.active)
    _build_field_sheet(wb, "Deal Setup", DEAL_FIELDS)
    _build_operating_sheet(wb, years)
    _build_debt_sheet(wb, years)
    _build_field_sheet(wb, "Return Assumptions", RETURN_FIELDS)
    _build_field_sheet(wb, "Qualitative Inputs", QUAL_FIELDS)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        for col in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col)].width = 22
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _build_field_sheet(wb: Workbook, title: str, rows: list[tuple[str, str, object]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(["field_id", "display_name", "value"])
    _style_header(ws)
    for field_id, display_name, value in rows:
        ws.append([field_id, display_name, value])


def _build_operating_sheet(wb: Workbook, years: int) -> None:
    ws = wb.create_sheet("Operating Forecast")
    ws.append(["year", "revenue", "ebitda", "ebitda_margin", "capex", "cash_tax", "change_in_nwc", "fcf"])
    _style_header(ws)
    start_year = 2027
    revenue = 500.0
    for idx in range(years):
        year = start_year + idx
        revenue *= 1.08
        ebitda = revenue * 0.22
        capex = revenue * 0.04
        cash_tax = ebitda * 0.20
        nwc = revenue * 0.01
        fcf = ebitda - capex - cash_tax - nwc
        ws.append([year, revenue, ebitda, ebitda / revenue, capex, cash_tax, nwc, fcf])


def _build_debt_sheet(wb: Workbook, years: int) -> None:
    ws = wb.create_sheet("Debt & Cash")
    ws.append(["year", "opening_debt", "new_debt", "repayment", "cash_sweep", "ending_debt", "cash", "net_debt"])
    _style_header(ws)
    debt = 400.0
    cash = 50.0
    for idx in range(years):
        year = 2027 + idx
        opening = debt
        repayment = 40.0
        debt = max(debt - repayment, 0.0)
        cash += 5.0
        ws.append([year, opening, 0.0, repayment, 0.0, debt, cash, debt - cash])


def _style_header(ws) -> None:
    fill = PatternFill("solid", fgColor="1F2937")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
