from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from openpyxl.workbook.workbook import Workbook


class WorkbookLoadError(RuntimeError):
    """Raised when an uploaded workbook cannot be opened as a normal xlsx file."""


WorkbookInput = str | Path | bytes | BinaryIO | BytesIO


class ExcelLoader:
    def __init__(self, workbook_file: WorkbookInput, data_only: bool = True):
        self.workbook_file = workbook_file
        self.data_only = data_only
        self.workbook_name = self._name_for(workbook_file)
        self._workbook: Workbook | None = None

    def __enter__(self) -> "ExcelLoader":
        self._workbook = self._load()
        return self

    def __exit__(self, exc_type: object, exc: object, traceback: object) -> None:
        if self._workbook is not None:
            self._workbook.close()

    @property
    def workbook(self) -> Workbook:
        if self._workbook is None:
            self._workbook = self._load()
        return self._workbook

    def list_sheets(self) -> list[str]:
        return list(self.workbook.sheetnames)

    def read_cell(self, sheet: str, cell: str) -> Any:
        return self.workbook[sheet][cell].value

    def read_range(self, sheet: str, range_ref: str) -> pd.DataFrame:
        worksheet = self.workbook[sheet]
        values = [[cell.value for cell in row] for row in worksheet[range_ref]]
        return pd.DataFrame(values)

    def read_flat_range(self, sheet: str, range_ref: str) -> list[Any]:
        worksheet = self.workbook[sheet]
        min_col, min_row, max_col, max_row = range_boundaries(range_ref)
        values: list[Any] = []
        for row in worksheet.iter_rows(
            min_row=min_row,
            max_row=max_row,
            min_col=min_col,
            max_col=max_col,
        ):
            values.extend(cell.value for cell in row)
        return values

    def _load(self) -> Workbook:
        try:
            source = self._source_for_openpyxl(self.workbook_file)
            return load_workbook(source, read_only=True, data_only=self.data_only)
        except Exception as exc:  # openpyxl exposes several low-level parse failures.
            raise WorkbookLoadError(f"Unable to load workbook '{self.workbook_name}': {exc}") from exc

    @staticmethod
    def _source_for_openpyxl(workbook_file: WorkbookInput) -> WorkbookInput:
        if isinstance(workbook_file, bytes):
            return BytesIO(workbook_file)
        if hasattr(workbook_file, "getvalue"):
            return BytesIO(workbook_file.getvalue())
        return workbook_file

    @staticmethod
    def _name_for(workbook_file: WorkbookInput) -> str:
        if isinstance(workbook_file, (str, Path)):
            return Path(workbook_file).name
        return getattr(workbook_file, "name", "uploaded_workbook.xlsx")

